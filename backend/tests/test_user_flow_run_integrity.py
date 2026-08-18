"""Unit tests for Phase UP3-1: Run integrity, lifecycle, publication contract, and orphan cleanup."""

from __future__ import annotations

import json
import pytest
from typing import Any

import uuid
from domain.user_flow import import_user_flow_xlsx, run_user_flow_alignment, UserFlowRunResult
from domain.user_flow._queries import get_user_flow_report
from domain.user_flow._extract import WorkbookExtract, SheetExtract, Block, RawRow, RawCell
from shared.utils import utc_now_iso


def new_id() -> str:
    return str(uuid.uuid4())


@pytest.fixture
def memory_db_schema():
    """Create in-memory SQLite schema for testing."""
    import aiosqlite
    from infrastructure.db.database.migrations import _MIGRATIONS

    async def _make_db():
        db = await aiosqlite.connect(":memory:")
        db.row_factory = aiosqlite.Row
        for m in _MIGRATIONS:
            await db.executescript(m["sql"])
        await db.commit()
        return db

    return _make_db


@pytest.mark.asyncio
async def test_two_run_staleness_and_completion_gating(memory_db_schema):
    """Run A completes; Run B fails midway -> report returns Run A's data intact with run_id=run_a."""
    db = await memory_db_schema()
    doc_id = f"ufdoc:{new_id()}"
    cluster_id = "test_cluster"
    snapshot_id = "test_snapshot"
    flow_id = f"uf:{new_id()}"
    step_id = f"ustep:{new_id()}"
    activity_id = f"uact:{new_id()}"
    now = utc_now_iso()

    # Seed baseline doc, flow, step, activity
    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES (?, 'test.xlsx', 'h1', ?)", (doc_id, now))
    await db.execute("INSERT INTO user_flows (id, doc_id, ordinal, name_ja, name_en, kind, sheet, scope_note, created_at) VALUES (?, ?, 1, 'Flow 1', 'Flow 1', 'narrative', 'Sheet1', '', ?)", (flow_id, doc_id, now))
    await db.execute("INSERT INTO user_steps (id, flow_id, ordinal, kind, section_id, text_ja, in_scope, sheet, row_start, row_end, created_at) VALUES (?, ?, 1, 'action', '1', 'Step 1', 1, 'Sheet1', 1, 2, ?)", (step_id, flow_id, now))
    await db.execute("INSERT INTO user_activities (id, flow_id, ordinal, name_en, name_ja, member_step_ids_json, sheet_span_json, step_count, origin, created_at) VALUES (?, ?, 1, 'Act 1', 'Act 1', ?, '[]', 1, 'llm', ?)", (activity_id, flow_id, json.dumps([step_id]), now))

    # Run A: Completed run
    run_a = "ufrun:run_a"
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, confidence, reason, created_at) VALUES ('m_a', ?, ?, 'bf1', 'FULLY', 0.95, 'Matched in run A', ?)", (run_a, activity_id, now))
    await db.execute("INSERT INTO user_bd_mappings (id, run_id, user_step_id, bd_kind, bd_id, relation, confidence, reason, created_at) VALUES ('map_a', ?, ?, 'step', 'bs1', 'realizes', 0.9, 'Mapped in run A', ?)", (run_a, step_id, now))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, reason, created_at) VALUES ('v_a', ?, ?, ?, ?, 'user', ?, 'step', 'COVERED', 'Verdict in run A', ?)", (run_a, doc_id, cluster_id, snapshot_id, step_id, now))
    await db.execute(
        "INSERT INTO user_run_artifacts (id, run_id, doc_id, ref_id, payload, created_at) VALUES ('art_a', ?, ?, '__run_complete__', ?, ?)",
        (run_a, doc_id, json.dumps({"doc_id": doc_id, "run_id": run_a, "cluster_id": cluster_id, "snapshot_id": snapshot_id, "status": "ok"}), now),
    )

    # Run B: In-progress or failed midway run (NO __run_complete__ artifact, has toxic NONE match and UNVERIFIABLE verdict)
    run_b = "ufrun:run_b"
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, confidence, reason, created_at) VALUES ('m_b', ?, ?, NULL, 'NONE', 0.1, 'Failed run B', ?)", (run_b, activity_id, now))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, reason, created_at) VALUES ('v_b', ?, ?, ?, ?, 'user', ?, 'step', 'UNVERIFIABLE', 'Incomplete run B', ?)", (run_b, doc_id, cluster_id, snapshot_id, step_id, now))
    await db.commit()

    # Query report without explicit run_id -> must resolve to Run A
    rep = await get_user_flow_report(db, doc_id, cluster_id, snapshot_id)
    assert rep["summary"]["run_id"] == run_a
    assert rep["run_id"] == run_a

    # Check that Run A's match and verdict are returned, not Run B's
    flow_rep = rep["flows"][0]
    act_rep = flow_rep["activities"][0]
    assert act_rep["match_status"] == "FULLY"
    assert act_rep["reason"] != "Failed run B"

    step_rep = flow_rep["steps"][0]
    assert step_rep["verdict"] == "COVERED"
    assert step_rep["reason"] == "Verdict in run A"

    # Pinning the failed run must NOT expose it: a pin narrows the search, it does not waive the
    # completion marker. Run B's toxic NONE/UNVERIFIABLE rows stay invisible.
    rep_b = await get_user_flow_report(db, doc_id, cluster_id, snapshot_id, run_id=run_b)
    assert rep_b["run_id"] is None
    assert rep_b["summary"]["run_id"] is None
    assert rep_b["flows"][0]["activities"][0]["pair_matches"] == []
    assert rep_b["flows"][0]["steps"][0]["reason"] == ""

    # Pinning the completed run is allowed and returns exactly that run.
    rep_a = await get_user_flow_report(db, doc_id, cluster_id, snapshot_id, run_id=run_a)
    assert rep_a["run_id"] == run_a
    assert rep_a["flows"][0]["steps"][0]["reason"] == "Verdict in run A"

    await db.close()


@pytest.mark.asyncio
async def test_report_never_serves_another_clusters_completed_run(memory_db_schema):
    """The marker records the (doc, cluster, snapshot) the run analysed; a newer run of a DIFFERENT
    cluster/snapshot must never answer this request, even though it shares the document."""
    db = await memory_db_schema()
    doc_id = f"ufdoc:{new_id()}"
    flow_id = f"uf:{new_id()}"
    step_id = f"ustep:{new_id()}"
    activity_id = f"uact:{new_id()}"
    now = utc_now_iso()

    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES (?, 'test.xlsx', 'h1', ?)", (doc_id, now))
    await db.execute("INSERT INTO user_flows (id, doc_id, ordinal, name_ja, name_en, kind, sheet, scope_note, created_at) VALUES (?, ?, 1, 'Flow 1', 'Flow 1', 'narrative', 'Sheet1', '', ?)", (flow_id, doc_id, now))
    await db.execute("INSERT INTO user_steps (id, flow_id, ordinal, kind, section_id, text_ja, in_scope, sheet, row_start, row_end, created_at) VALUES (?, ?, 1, 'action', '1', 'Step 1', 1, 'Sheet1', 1, 2, ?)", (step_id, flow_id, now))
    await db.execute("INSERT INTO user_activities (id, flow_id, ordinal, name_en, name_ja, member_step_ids_json, sheet_span_json, step_count, origin, created_at) VALUES (?, ?, 1, 'Act 1', 'Act 1', ?, '[]', 1, 'llm', ?)", (activity_id, flow_id, json.dumps([step_id]), now))

    # Older completed run for (c1, sn1) — the one this request must get.
    run_1 = "ufrun:c1"
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, confidence, reason, created_at) VALUES ('m1', ?, ?, 'bf1', 'FULLY', 0.9, 'c1 run', ?)", (run_1, activity_id, now))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, reason, created_at) VALUES ('v1', ?, ?, 'c1', 'sn1', 'user', ?, 'step', 'COVERED', 'c1 verdict', ?)", (run_1, doc_id, step_id, now))
    await db.execute(
        "INSERT INTO user_run_artifacts (id, run_id, doc_id, ref_id, payload, created_at) VALUES ('a1', ?, ?, '__run_complete__', ?, '2026-08-01T10:00:00+00:00')",
        (run_1, doc_id, json.dumps({"doc_id": doc_id, "run_id": run_1, "cluster_id": "c1", "snapshot_id": "sn1"})),
    )

    # NEWER completed run for a different (c2, sn2).
    run_2 = "ufrun:c2"
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, confidence, reason, created_at) VALUES ('m2', ?, ?, NULL, 'NONE', 0.9, 'c2 run', ?)", (run_2, activity_id, now))
    await db.execute(
        "INSERT INTO user_run_artifacts (id, run_id, doc_id, ref_id, payload, created_at) VALUES ('a2', ?, ?, '__run_complete__', ?, '2026-08-09T10:00:00+00:00')",
        (run_2, doc_id, json.dumps({"doc_id": doc_id, "run_id": run_2, "cluster_id": "c2", "snapshot_id": "sn2"})),
    )
    await db.commit()

    rep = await get_user_flow_report(db, doc_id, "c1", "sn1")
    assert rep["run_id"] == run_1
    assert rep["flows"][0]["activities"][0]["match_status"] == "FULLY"
    assert rep["flows"][0]["steps"][0]["reason"] == "c1 verdict"

    # Pinning the other cluster's run from a (c1, sn1) request resolves to nothing.
    rep_x = await get_user_flow_report(db, doc_id, "c1", "sn1", run_id=run_2)
    assert rep_x["run_id"] is None

    await db.close()


@pytest.mark.asyncio
async def test_report_never_selects_a_run_without_completion_marker(memory_db_schema):
    """Publication contract: a run that never wrote __run_complete__ is invisible to the report."""
    db = await memory_db_schema()
    doc_id = f"ufdoc:{new_id()}"
    cluster_id = "test_cluster"
    snapshot_id = "test_snapshot"
    flow_id = f"uf:{new_id()}"
    step_id = f"ustep:{new_id()}"
    activity_id = f"uact:{new_id()}"
    now = utc_now_iso()

    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES (?, 'test.xlsx', 'h1', ?)", (doc_id, now))
    await db.execute("INSERT INTO user_flows (id, doc_id, ordinal, name_ja, name_en, kind, sheet, scope_note, created_at) VALUES (?, ?, 1, 'Flow 1', 'Flow 1', 'narrative', 'Sheet1', '', ?)", (flow_id, doc_id, now))
    await db.execute("INSERT INTO user_steps (id, flow_id, ordinal, kind, section_id, text_ja, in_scope, sheet, row_start, row_end, created_at) VALUES (?, ?, 1, 'action', '1', 'Step 1', 1, 'Sheet1', 1, 2, ?)", (step_id, flow_id, now))
    await db.execute("INSERT INTO user_activities (id, flow_id, ordinal, name_en, name_ja, member_step_ids_json, sheet_span_json, step_count, origin, created_at) VALUES (?, ?, 1, 'Act 1', 'Act 1', ?, '[]', 1, 'llm', ?)", (activity_id, flow_id, json.dumps([step_id]), now))

    # The only run present never completed: all-NONE tier-1 rows + a UNVERIFIABLE verdict.
    run_x = "ufrun:never_completed"
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, confidence, reason, created_at) VALUES ('m_x', ?, ?, NULL, 'NONE', 0.1, 'LLM_NO_RESPONSE', ?)", (run_x, activity_id, now))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, reason, created_at) VALUES ('v_x', ?, ?, ?, ?, 'user', ?, 'step', 'UNVERIFIABLE', 'partial run', ?)", (run_x, doc_id, cluster_id, snapshot_id, step_id, now))
    await db.commit()

    rep = await get_user_flow_report(db, doc_id, cluster_id, snapshot_id)
    assert rep["run_id"] is None
    assert rep["summary"]["run_id"] is None
    # None of the incomplete run's rows may surface.
    act_rep = rep["flows"][0]["activities"][0]
    assert act_rep["pair_matches"] == []
    assert act_rep["match_status"] == "UNRESOLVED"  # never "NONE": nothing was ever published
    assert rep["flows"][0]["steps"][0]["reason"] == ""

    # A pin narrows the search; it never waives the completion marker.
    rep_pinned = await get_user_flow_report(db, doc_id, cluster_id, snapshot_id, run_id=run_x)
    assert rep_pinned["run_id"] is None
    assert rep_pinned["flows"][0]["steps"][0]["reason"] == ""

    await db.close()


@pytest.mark.asyncio
async def test_mid_run_failure_leaves_previous_report_intact(memory_db_schema, tmp_path, monkeypatch):
    """A run that raises inside the verdict stage publishes nothing: report still serves run A."""
    import domain.user_flow as user_flow_pkg

    db = await memory_db_schema()
    doc_id = f"ufdoc:{new_id()}"
    cluster_id = "test_cluster"
    snapshot_id = f"snap:{new_id()}"
    flow_id = f"uf:{new_id()}"
    step_id = f"ustep:{new_id()}"
    activity_id = f"uact:{new_id()}"
    bd_flow_id = f"bf:{new_id()}"
    now = utc_now_iso()

    await db.execute("INSERT INTO repo_snapshots (id, local_repo_id, local_path, synced_at, created_at) VALUES (?, 'repo', ?, ?, ?)", (snapshot_id, str(tmp_path), now, now))
    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES (?, 'test.xlsx', 'h1', ?)", (doc_id, now))
    await db.execute("INSERT INTO user_flows (id, doc_id, ordinal, name_ja, name_en, kind, sheet, scope_note, created_at) VALUES (?, ?, 1, 'Flow 1', 'Flow 1', 'narrative', 'Sheet1', '', ?)", (flow_id, doc_id, now))
    await db.execute("INSERT INTO user_steps (id, flow_id, ordinal, kind, section_id, text_ja, in_scope, sheet, row_start, row_end, created_at) VALUES (?, ?, 1, 'action', '1', 'Step 1', 1, 'Sheet1', 1, 2, ?)", (step_id, flow_id, now))
    await db.execute("INSERT INTO user_activities (id, flow_id, ordinal, name_en, name_ja, member_step_ids_json, sheet_span_json, step_count, origin, created_at) VALUES (?, ?, 1, 'Act 1', 'Act 1', ?, '[]', 1, 'llm', ?)", (activity_id, flow_id, json.dumps([step_id]), now))
    await db.execute("INSERT INTO bd_business_flows (id, cluster_id, doc_id, sub_ix, block_key, name, description, ordinal, origin, created_at) VALUES (?, ?, 'doc1', 1, 'blk1', 'BD Flow', 'Desc', 1, 'declared', ?)", (bd_flow_id, cluster_id, now))
    await db.commit()

    # Run A completes normally (offline), then its tier-1 row is upgraded to a real FULLY match.
    res_a = await run_user_flow_alignment(db, doc_id, cluster_id, snapshot_id, provider_id=None)
    await db.execute(
        "UPDATE user_activity_matches SET match_status = 'FULLY', bd_flow_id = ?, reason = 'good run A' WHERE run_id = ?",
        (bd_flow_id, res_a.run_id),
    )
    await db.commit()

    # Run B blows up in the verdict stage after tier-1 already wrote its all-NONE rows.
    def _boom(*args: Any, **kwargs: Any):
        raise RuntimeError("verdict stage crashed")

    monkeypatch.setattr(user_flow_pkg, "run_user_flow_verdicts", _boom)
    with pytest.raises(RuntimeError):
        await run_user_flow_alignment(db, doc_id, cluster_id, snapshot_id, provider_id=None)

    # Run B wrote tier-1 rows but no completion marker.
    async with db.execute("SELECT COUNT(*) FROM user_activity_matches WHERE run_id != ?", (res_a.run_id,)) as cur:
        assert (await cur.fetchone())[0] > 0
    async with db.execute("SELECT COUNT(*) FROM user_run_artifacts WHERE doc_id = ? AND ref_id = '__run_complete__'", (doc_id,)) as cur:
        assert (await cur.fetchone())[0] == 1

    rep = await get_user_flow_report(db, doc_id, cluster_id, snapshot_id)
    assert rep["run_id"] == res_a.run_id
    act_rep = rep["flows"][0]["activities"][0]
    assert act_rep["match_status"] == "FULLY"
    assert [p["reason"] for p in act_rep["pair_matches"]] == ["good run A"]

    await db.close()


@pytest.mark.asyncio
async def test_reimport_leaves_zero_orphans(memory_db_schema, tmp_path):
    """Re-importing a doc must delete all dependent alignment rows cleanly with zero orphans."""
    db = await memory_db_schema()
    cluster_id = "test_cluster"
    snapshot_id = "test_snapshot"

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="No.")
    ws.cell(row=1, column=2, value="操作内容")
    ws.cell(row=2, column=1, value="1")
    ws.cell(row=2, column=2, value="Step 1")
    xlsx_path = tmp_path / "test.xlsx"
    wb.save(xlsx_path)

    # 1. Initial import
    doc_id1 = await import_user_flow_xlsx(db, xlsx_path, provider_id=None)

    # Seed structured flow, step, and activity for doc_id1
    now = utc_now_iso()
    flow_id = f"uf:{new_id()}"
    step_id = f"ustep:{new_id()}"
    activity_id = f"uact:{new_id()}"
    await db.execute("INSERT INTO user_flows (id, doc_id, ordinal, name_ja, name_en, kind, sheet, scope_note, created_at) VALUES (?, ?, 1, 'Flow 1', 'Flow 1', 'narrative', 'Sheet1', '', ?)", (flow_id, doc_id1, now))
    await db.execute("INSERT INTO user_steps (id, flow_id, ordinal, kind, section_id, text_ja, in_scope, sheet, row_start, row_end, created_at) VALUES (?, ?, 1, 'action', '1', 'Step 1', 1, 'Sheet1', 1, 2, ?)", (step_id, flow_id, now))
    await db.execute("INSERT INTO user_activities (id, flow_id, ordinal, name_en, name_ja, member_step_ids_json, sheet_span_json, step_count, origin, created_at) VALUES (?, ?, 1, 'Act 1', 'Act 1', ?, '[]', 1, 'llm', ?)", (activity_id, flow_id, json.dumps([step_id]), now))

    # Seed alignment rows on doc_id1
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, created_at) VALUES ('m1', 'r1', ?, 'bf1', 'FULLY', ?)", (activity_id, now))
    await db.execute("INSERT INTO user_bd_mappings (id, run_id, user_step_id, bd_kind, bd_id, relation, created_at) VALUES ('map1', 'r1', ?, 'step', 'bs1', 'realizes', ?)", (step_id, now))
    await db.execute("INSERT INTO user_code_anchors (id, run_id, snapshot_id, step_id, rel_path, line_start, line_end, kind, valid, created_at) VALUES ('anc1', 'r1', ?, ?, 'foo.cbl', 1, 5, 'llm_matched', 1, ?)", (snapshot_id, step_id, now))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, created_at) VALUES ('v1', 'r1', ?, ?, ?, 'user', ?, 'step', 'COVERED', ?)", (doc_id1, cluster_id, snapshot_id, step_id, now))
    await db.commit()

    # 2. Re-import same workbook (same file hash triggers idempotency cleanup)
    doc_id2 = await import_user_flow_xlsx(db, xlsx_path, provider_id=None)
    assert doc_id2 != doc_id1

    # Verify that NO orphan rows reference old ids
    async with db.execute("SELECT COUNT(*) FROM user_bd_mappings WHERE user_step_id = ?", (step_id,)) as cur:
        assert (await cur.fetchone())[0] == 0
    async with db.execute("SELECT COUNT(*) FROM user_code_anchors WHERE step_id = ?", (step_id,)) as cur:
        assert (await cur.fetchone())[0] == 0
    async with db.execute("SELECT COUNT(*) FROM user_verdicts WHERE doc_id = ?", (doc_id1,)) as cur:
        assert (await cur.fetchone())[0] == 0
    async with db.execute("SELECT COUNT(*) FROM user_activity_matches WHERE activity_id = ?", (activity_id,)) as cur:
        assert (await cur.fetchone())[0] == 0
    async with db.execute("SELECT COUNT(*) FROM user_activities WHERE flow_id = ?", (flow_id,)) as cur:
        assert (await cur.fetchone())[0] == 0

    await db.close()


@pytest.mark.asyncio
async def test_garbage_collection_keeps_last_two_completed_runs(memory_db_schema):
    """When a 3rd run completes, the oldest run is safely pruned, retaining current + 1 prior."""
    db = await memory_db_schema()
    doc_id = f"ufdoc:{new_id()}"
    cluster_id = "test_cluster"
    snapshot_id = "test_snapshot"
    flow_id = f"uf:{new_id()}"
    step_id = f"ustep:{new_id()}"
    activity_id = f"uact:{new_id()}"
    now = utc_now_iso()

    await db.execute("INSERT INTO repo_snapshots (id, local_repo_id, local_path, synced_at, created_at) VALUES (?, 'repo', 'd:/dummy', ?, ?)", (snapshot_id, now, now))
    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES (?, 'test.xlsx', 'h1', ?)", (doc_id, now))
    await db.execute("INSERT INTO user_flows (id, doc_id, ordinal, name_ja, name_en, kind, sheet, scope_note, created_at) VALUES (?, ?, 1, 'Flow 1', 'Flow 1', 'narrative', 'Sheet1', '', ?)", (flow_id, doc_id, now))
    await db.execute("INSERT INTO user_steps (id, flow_id, ordinal, kind, section_id, text_ja, in_scope, sheet, row_start, row_end, created_at) VALUES (?, ?, 1, 'action', '1', 'Step 1', 1, 'Sheet1', 1, 2, ?)", (step_id, flow_id, now))
    await db.execute("INSERT INTO user_activities (id, flow_id, ordinal, name_en, name_ja, member_step_ids_json, sheet_span_json, step_count, origin, created_at) VALUES (?, ?, 1, 'Act 1', 'Act 1', ?, '[]', 1, 'llm', ?)", (activity_id, flow_id, json.dumps([step_id]), now))
    await db.commit()

    def _marker(cl: str, sn: str, rid: str) -> str:
        return json.dumps({"doc_id": doc_id, "run_id": rid, "cluster_id": cl, "snapshot_id": sn})

    # Simulate Run 1 (oldest completed, same tuple)
    run1 = "ufrun:run1"
    await db.execute("INSERT INTO user_run_artifacts (id, run_id, doc_id, ref_id, payload, created_at) VALUES ('art_c1', ?, ?, '__run_complete__', ?, '2026-08-01T10:00:00+00:00')", (run1, doc_id, _marker(cluster_id, snapshot_id, run1)))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, created_at) VALUES ('v_r1', ?, ?, ?, ?, 'user', ?, 'step', 'COVERED', '2026-08-01T10:00:00+00:00')", (run1, doc_id, cluster_id, snapshot_id, step_id))

    # Simulate Run 2 (prior completed, same tuple)
    run2 = "ufrun:run2"
    await db.execute("INSERT INTO user_run_artifacts (id, run_id, doc_id, ref_id, payload, created_at) VALUES ('art_c2', ?, ?, '__run_complete__', ?, '2026-08-02T11:00:00+00:00')", (run2, doc_id, _marker(cluster_id, snapshot_id, run2)))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, created_at) VALUES ('v_r2', ?, ?, ?, ?, 'user', ?, 'step', 'COVERED', '2026-08-02T11:00:00+00:00')", (run2, doc_id, cluster_id, snapshot_id, step_id))

    # An OLD completed run of a DIFFERENT cluster/snapshot: outside this run's GC scope.
    run_other = "ufrun:other_cluster"
    await db.execute("INSERT INTO user_run_artifacts (id, run_id, doc_id, ref_id, payload, created_at) VALUES ('art_co', ?, ?, '__run_complete__', ?, '2026-07-01T10:00:00+00:00')", (run_other, doc_id, _marker("other_cluster", "other_snapshot", run_other)))
    await db.execute("INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, created_at) VALUES ('v_ro', ?, ?, 'other_cluster', 'other_snapshot', 'user', ?, 'step', 'COVERED', '2026-07-01T10:00:00+00:00')", (run_other, doc_id, step_id))

    # An UNMARKED run (possibly still in flight): never collected by completion GC.
    run_inflight = "ufrun:inflight"
    await db.execute("INSERT INTO user_run_artifacts (id, run_id, doc_id, ref_id, payload, created_at) VALUES ('art_if', ?, ?, ?, '{}', ?)", (run_inflight, doc_id, f"align:{step_id}", now))
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, reason, created_at) VALUES ('am_if', ?, ?, NULL, 'NONE', 'in flight', ?)", (run_inflight, activity_id, now))
    await db.commit()

    # Execute Run 3 via run_user_flow_alignment (offline mode, provider_id=None)
    run_res = await run_user_flow_alignment(db, doc_id, cluster_id, snapshot_id, provider_id=None)
    run3 = run_res.run_id

    # Run 1 (3rd-newest completed of this tuple) is pruned; run 2 + run 3 stay.
    async with db.execute("SELECT DISTINCT run_id FROM user_verdicts WHERE doc_id = ?", (doc_id,)) as cur:
        verdict_runs = {r[0] for r in await cur.fetchall()}
    assert run1 not in verdict_runs
    assert run2 in verdict_runs
    assert run3 in verdict_runs
    assert run_other in verdict_runs, "GC crossed into another cluster/snapshot's completed run"

    async with db.execute("SELECT DISTINCT run_id FROM user_run_artifacts WHERE doc_id = ? AND ref_id = '__run_complete__'", (doc_id,)) as cur:
        artifact_runs = {r[0] for r in await cur.fetchall()}
    assert run1 not in artifact_runs
    assert run2 in artifact_runs
    assert run3 in artifact_runs
    assert run_other in artifact_runs

    # The unmarked run is untouched — it may still be writing.
    async with db.execute("SELECT COUNT(*) FROM user_run_artifacts WHERE run_id = ?", (run_inflight,)) as cur:
        assert (await cur.fetchone())[0] == 1
    async with db.execute("SELECT COUNT(*) FROM user_activity_matches WHERE run_id = ?", (run_inflight,)) as cur:
        assert (await cur.fetchone())[0] == 1

    await db.close()


@pytest.mark.asyncio
async def test_failed_finalization_leaves_no_pending_marker(memory_db_schema, monkeypatch):
    """The DB handle is a process-wide singleton: a marker written but not committed would be
    published by the NEXT unrelated commit. Finalization must roll back as a unit."""
    from domain.user_flow import _publish_completed_run

    db = await memory_db_schema()
    doc_id = f"ufdoc:{new_id()}"
    cluster_id, snapshot_id, run_id = "cl", "sn", "ufrun:boom"
    now = utc_now_iso()
    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES (?, 'x.xlsx', 'h', ?)", (doc_id, now))
    await db.commit()

    real_execute = db.execute
    calls = {"n": 0}

    def _exploding_execute(sql: str, *args: Any, **kwargs: Any):
        # Fail the GC scan that runs right after the marker insert.
        if "__run_complete__" in sql and sql.strip().upper().startswith("SELECT"):
            calls["n"] += 1
            raise RuntimeError("connection dropped mid-finalization")
        return real_execute(sql, *args, **kwargs)

    monkeypatch.setattr(db, "execute", _exploding_execute)
    with pytest.raises(RuntimeError):
        await _publish_completed_run(
            db, doc_id=doc_id, cluster_id=cluster_id, snapshot_id=snapshot_id,
            run_id=run_id, summary={"doc_id": doc_id, "run_id": run_id, "cluster_id": cluster_id, "snapshot_id": snapshot_id},
        )
    assert calls["n"] == 1
    monkeypatch.undo()

    # A later, unrelated commit must not publish the aborted run.
    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES ('ufdoc:other', 'y.xlsx', 'h2', ?)", (now,))
    await db.commit()

    async with db.execute("SELECT COUNT(*) FROM user_run_artifacts WHERE ref_id = '__run_complete__'") as cur:
        assert (await cur.fetchone())[0] == 0
    assert await get_user_flow_report(db, doc_id, cluster_id, snapshot_id) is not None
    rep = await get_user_flow_report(db, doc_id, cluster_id, snapshot_id)
    assert rep["run_id"] is None

    await db.close()


@pytest.mark.asyncio
async def test_orphan_prune_migration_keeps_live_rows():
    """The one-time prune (migration 17) removes dead-id rows only — BD-side verdicts, whose
    ref_kind vocabulary overlaps the user side, must survive."""
    import aiosqlite
    from infrastructure.db.database.migrations import _MIGRATIONS

    db = await aiosqlite.connect(":memory:")
    db.row_factory = aiosqlite.Row
    prune = next(m for m in _MIGRATIONS if m["version"] == 17)
    for m in _MIGRATIONS:
        if m["version"] != 17:
            await db.executescript(m["sql"])

    now = utc_now_iso()
    doc_id, flow_id, step_id, act_id = "ufdoc:live", "uf:live", "ustep:live", "uact:live"
    await db.execute("INSERT INTO user_flow_docs (id, source_name, file_hash, imported_at) VALUES (?, 'x.xlsx', 'h', ?)", (doc_id, now))
    await db.execute("INSERT INTO user_flows (id, doc_id, ordinal, name_ja, name_en, kind, sheet, scope_note, created_at) VALUES (?, ?, 1, 'F', 'F', 'n', 'S', '', ?)", (flow_id, doc_id, now))
    await db.execute("INSERT INTO user_steps (id, flow_id, ordinal, kind, section_id, text_ja, in_scope, sheet, row_start, row_end, created_at) VALUES (?, ?, 1, 'action', '1', 'S', 1, 'S', 1, 2, ?)", (step_id, flow_id, now))
    await db.execute("INSERT INTO user_activities (id, flow_id, ordinal, name_en, name_ja, member_step_ids_json, sheet_span_json, step_count, origin, created_at) VALUES (?, ?, 1, 'A', 'A', ?, '[]', 1, 'llm', ?)", (act_id, flow_id, json.dumps([step_id]), now))

    live_rows = [
        ("v_user_step", "user", step_id, "step"),
        ("v_user_flow", "user", flow_id, "flow"),
        ("v_user_act", "user", act_id, "activity"),
        ("v_bd_step", "bd", "bd_step_1", "step"),      # BD unit id, NOT a user_steps id
        ("v_bd_flow", "bd", "bf1", "flow"),
        ("v_bd_branch", "bd", "bb1", "branch"),
    ]
    for vid, side, ref_id, ref_kind in live_rows:
        await db.execute(
            "INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, created_at) "
            "VALUES (?, 'r1', ?, 'cl', 'snap', ?, ?, ?, 'COVERED', ?)",
            (vid, doc_id, side, ref_id, ref_kind, now),
        )
    # Genuine orphans: dead step id, and a whole dead document.
    await db.execute(
        "INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, created_at) "
        "VALUES ('v_orphan_step', 'r0', ?, 'cl', 'snap', 'user', 'ustep:dead', 'step', 'COVERED', ?)",
        (doc_id, now),
    )
    await db.execute(
        "INSERT INTO user_verdicts (id, run_id, doc_id, cluster_id, snapshot_id, side, ref_id, ref_kind, verdict, created_at) "
        "VALUES ('v_orphan_doc', 'r0', 'ufdoc:dead', 'cl', 'snap', 'bd', 'bd_step_9', 'step', 'BD_EXTRA', ?)",
        (now,),
    )
    await db.execute("INSERT INTO user_bd_mappings (id, run_id, user_step_id, bd_kind, bd_id, relation, created_at) VALUES ('m_live', 'r1', ?, 'step', 'bs1', 'realizes', ?)", (step_id, now))
    await db.execute("INSERT INTO user_bd_mappings (id, run_id, user_step_id, bd_kind, bd_id, relation, created_at) VALUES ('m_orphan', 'r0', 'ustep:dead', 'step', 'bs1', 'realizes', ?)", (now,))
    await db.execute("INSERT INTO user_code_anchors (id, run_id, snapshot_id, step_id, rel_path, line_start, line_end, kind, valid, created_at) VALUES ('a_orphan', 'r0', 'snap', 'ustep:dead', 'f.cbl', 1, 2, 'llm_matched', 1, ?)", (now,))
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, reason, created_at) VALUES ('am_live', 'r1', ?, 'bf1', 'FULLY', 'ok', ?)", (act_id, now))
    await db.execute("INSERT INTO user_activity_matches (id, run_id, activity_id, bd_flow_id, match_status, reason, created_at) VALUES ('am_failed', 'r_failed', ?, NULL, 'NONE', 'LLM_NO_RESPONSE', ?)", (act_id, now))
    await db.commit()

    await db.executescript(prune["sql"])
    await db.commit()

    async with db.execute("SELECT id FROM user_verdicts ORDER BY id") as cur:
        surviving = {r[0] for r in await cur.fetchall()}
    assert surviving == {vid for vid, *_ in live_rows}, "prune removed live verdict rows"

    async with db.execute("SELECT id FROM user_bd_mappings") as cur:
        assert {r[0] for r in await cur.fetchall()} == {"m_live"}
    async with db.execute("SELECT COUNT(*) FROM user_code_anchors") as cur:
        assert (await cur.fetchone())[0] == 0
    async with db.execute("SELECT id FROM user_activity_matches") as cur:
        assert {r[0] for r in await cur.fetchall()} == {"am_live"}

    await db.close()


@pytest.mark.asyncio
async def test_gc_preserves_import_stage_artifacts(memory_db_schema, tmp_path):
    """Import artifacts are stored with run_id = doc_id; the run GC must never collect them."""
    db = await memory_db_schema()
    cluster_id = "test_cluster"
    snapshot_id = f"snap:{new_id()}"
    now = utc_now_iso()

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="No.")
    ws.cell(row=1, column=2, value="操作内容")
    ws.cell(row=2, column=1, value="1")
    ws.cell(row=2, column=2, value="Step 1")
    xlsx_path = tmp_path / "gc.xlsx"
    wb.save(xlsx_path)

    await db.execute("INSERT INTO repo_snapshots (id, local_repo_id, local_path, synced_at, created_at) VALUES (?, 'repo', ?, ?, ?)", (snapshot_id, str(tmp_path), now, now))
    await db.commit()

    doc_id = await import_user_flow_xlsx(db, xlsx_path, provider_id=None)

    async with db.execute("SELECT COUNT(*) FROM user_run_artifacts WHERE doc_id = ? AND ref_id = '__import_summary__'", (doc_id,)) as cur:
        assert (await cur.fetchone())[0] == 1

    # Two consecutive alignment runs exercise the GC path.
    await run_user_flow_alignment(db, doc_id, cluster_id, snapshot_id, provider_id=None)
    await run_user_flow_alignment(db, doc_id, cluster_id, snapshot_id, provider_id=None)

    async with db.execute("SELECT COUNT(*) FROM user_run_artifacts WHERE doc_id = ? AND ref_id = '__import_summary__'", (doc_id,)) as cur:
        assert (await cur.fetchone())[0] == 1, "run GC deleted the import-stage artifacts"

    await db.close()
